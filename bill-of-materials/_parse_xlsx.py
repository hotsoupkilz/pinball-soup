#!/usr/bin/env python3
"""Parse the Ozzfest BOM xlsx and write contents to a text file."""
import subprocess, sys

try:
    import openpyxl
except ImportError:
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'openpyxl'])
    import openpyxl

wb = openpyxl.load_workbook('/home/hotsoup/Downloads/Bill of Materials_ Ozzfest.xlsx')
output_path = '/home/hotsoup/repos/pinball-soup/bill-of-materials/_xlsx_parsed.txt'

with open(output_path, 'w') as f:
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        f.write(f'=== SHEET: {sheet_name} ({ws.max_row} rows x {ws.max_column} cols) ===\n')
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=False):
            vals = [str(cell.value) if cell.value is not None else '' for cell in row]
            f.write('\t'.join(vals) + '\n')
        f.write('\n')

print(f'Written to {output_path}')
